"""Create five polished Word questionnaires from the masked Markdown packets."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "user_study_materials"
OUTPUT_DIR = ROOT / "output" / "user_study_questionnaires"

NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float = 10.5, bold: bool | None = None, color: RGBColor = BLACK) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def add_inline_markdown(paragraph, text: str, size: float = 10.5) -> None:
    text = text.replace("  ", " ").strip()
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=bold)


def style_paragraph(paragraph, after: float = 5, before: float = 0, line: float = 1.08) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 16, 12, 7),
        ("Heading 2", 13, 10, 6),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(doc: Document, participant: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"AccountLens blinded evaluation questionnaire | Participant {participant}")
    set_run_font(run, size=8, color=RGBColor(89, 89, 89))


def add_cover(doc: Document, participant: str) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("AccountLens Blinded Evaluation Questionnaire")

    subtitle = doc.add_paragraph()
    run = subtitle.add_run(f"Participant {participant}   真人盲测问卷")
    set_run_font(run, size=14, bold=True)
    style_paragraph(subtitle, after=14)

    p = doc.add_paragraph()
    add_inline_markdown(
        p,
        "Purpose / 目的：评估两种账户简报在会议准备中的实用性、可核查性和清晰度。测试对象是简报，不是参与者。",
    )

    p = doc.add_paragraph()
    add_inline_markdown(
        p,
        "Time / 时长：约25至35分钟。所有公司、联系人和活动均为虚构数据。无需安装软件，也无需提供姓名、公司或客户资料。",
    )

    doc.add_heading("Consent and participant profile", level=1)
    consent_items = [
        "[ ] 我已了解测试目的，知道输出可能存在错误并需要人工确认。",
        "[ ] 我知道参与是自愿的，可以随时停止。",
        "[ ] 我同意仅以匿名编号记录时间、评分和意见。",
        "[ ] 我已年满18岁并同意参加本次测试。",
    ]
    for item in consent_items:
        p = doc.add_paragraph()
        add_inline_markdown(p, item)

    doc.add_paragraph("Background / 相关背景：")
    p = doc.add_paragraph()
    add_inline_markdown(
        p,
        "[ ] Account management   [ ] Sales   [ ] Presales   [ ] Customer success   [ ] Business analysis   [ ] Other",
    )
    p = doc.add_paragraph()
    add_inline_markdown(p, "Relevant experience / 相关经验：[ ] <1 year   [ ] 1-3 years   [ ] 4-7 years   [ ] 8+ years")

    doc.add_heading("What you need to complete", level=1)
    tasks = [
        "Review four fictional accounts. Each account contains Briefing A and Briefing B.",
        "Follow the stated review order and do not read the second briefing early.",
        "For each briefing, identify likely decision roles, two important risks, one next action and uncertain information.",
        "Record the number of seconds until you feel ready for the customer meeting.",
        "Score six dimensions from 1 to 5, where 1 is very poor and 5 is excellent.",
        "After each pair, choose A, B or tie and explain any error, unsupported claim or confusing abstention.",
    ]
    for task in tasks:
        p = doc.add_paragraph(style="List Bullet")
        add_inline_markdown(p, task)

    doc.add_heading("Rating scale", level=1)
    p = doc.add_paragraph()
    add_inline_markdown(p, "1 Very poor   2 Poor   3 Acceptable   4 Good   5 Excellent")
    p = doc.add_paragraph()
    add_inline_markdown(p, "The facilitator starts and stops the timer. Do not discuss your scores with other participants until the session ends.")

    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    headers = rows[0]
    if "Contact" in headers:
        widths = [1.15, 1.2, 1.0, 0.65, 1.55, 1.05]
        font_size = 7.7
    elif "Event ID" in headers:
        widths = [0.8, 0.7, 0.65, 1.0, 2.55, 0.8]
        font_size = 7.4
    else:
        widths = [7.15 / len(headers)] * len(headers)
        font_size = 8.5

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            repeat_table_header(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index in {2, 3, 5} else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(value)
            set_run_font(
                run,
                size=font_size,
                bold=row_index == 0,
                color=RGBColor(255, 255, 255) if row_index == 0 else BLACK,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_answer_lines(doc: Document, count: int = 2) -> None:
    for _ in range(count):
        p = doc.add_paragraph("________________________________________________________________________________")
        set_run_font(p.runs[0], size=9, color=RGBColor(128, 128, 128))
        style_paragraph(p, after=3)


def add_response_form(doc: Document, label: str, include_pair: bool) -> None:
    doc.add_page_break()
    doc.add_heading(f"Your response for Briefing {label}", level=2)
    p = doc.add_paragraph()
    add_inline_markdown(p, "请独立完成以下内容。主持人应在你表示已经准备好时停止计时。")

    prompts = [
        "1  Likely decision roles / 可能的决策角色",
        "2  Two most important risks / 两个最重要的风险",
        "3  Recommended next action / 建议的下一步行动",
        "4  Information that remains uncertain / 仍然不确定的信息",
    ]
    for prompt in prompts:
        p = doc.add_paragraph()
        run = p.add_run(prompt)
        set_run_font(run, size=10.5, bold=True)
        style_paragraph(p, before=4, after=2)
        add_answer_lines(doc, count=2)

    p = doc.add_paragraph()
    add_inline_markdown(p, "Meeting readiness / 会议准备状态：[ ] Ready   [ ] Partly ready   [ ] Not ready")
    p = doc.add_paragraph()
    add_inline_markdown(p, f"Briefing {label} preparation time / 准备时间：________ seconds")

    doc.add_heading("Ratings", level=3)
    rating_rows = [
        ["Measure / 评分项", "What to judge / 判断内容", "Score / 分数"],
        ["Role usefulness", "Does the role map support meeting preparation?", "[ ]1 [ ]2 [ ]3 [ ]4 [ ]5"],
        ["Evidence trust", "Can important claims be checked from event IDs?", "[ ]1 [ ]2 [ ]3 [ ]4 [ ]5"],
        ["Risk relevance", "Are the risks relevant to the next meeting?", "[ ]1 [ ]2 [ ]3 [ ]4 [ ]5"],
        ["Action usefulness", "Are the one or two actions practical?", "[ ]1 [ ]2 [ ]3 [ ]4 [ ]5"],
        ["Clarity", "Is the briefing clear and concise?", "[ ]1 [ ]2 [ ]3 [ ]4 [ ]5"],
        ["Uncertainty", "Is unknown or uncertain information clear?", "[ ]1 [ ]2 [ ]3 [ ]4 [ ]5"],
    ]
    add_table(doc, rating_rows)

    if include_pair:
        doc.add_heading("Paired assessment", level=2)
        p = doc.add_paragraph()
        add_inline_markdown(p, "Preferred briefing / 更偏好的简报：[ ] A   [ ] B   [ ] Tie")
        p = doc.add_paragraph()
        run = p.add_run("Why did you prefer it? / 选择原因")
        set_run_font(run, bold=True)
        add_answer_lines(doc, count=3)
        p = doc.add_paragraph()
        run = p.add_run("Errors, unsupported claims or confusing abstentions / 错误、缺少证据或难以理解的拒答")
        set_run_font(run, bold=True)
        add_answer_lines(doc, count=3)


def parse_markdown_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r"[-:]+", value) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def add_packet_content(doc: Document, packet_path: Path) -> None:
    lines = packet_path.read_text(encoding="utf-8").splitlines()
    # The generated Word cover replaces the short Markdown preamble.
    # The cover already ends with a page break, so skip the packet's first separator.
    index = next(i for i, line in enumerate(lines) if line.strip() == "---") + 1
    current_label = ""
    briefing_count_in_account = 0
    skip_generated_record_lines = False

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if skip_generated_record_lines:
            if line.startswith("### Briefing") or line == "---" or line.startswith("## Account"):
                skip_generated_record_lines = False
            else:
                index += 1
                continue

        if not line:
            index += 1
            continue
        if line == "---":
            doc.add_page_break()
            briefing_count_in_account = 0
            index += 1
            continue
        if line.startswith("## Account"):
            heading = doc.add_heading(line[3:].strip(), level=1)
            heading.paragraph_format.keep_with_next = True
            index += 1
            continue
        if line.startswith("**Review order:**"):
            p = doc.add_paragraph()
            add_inline_markdown(p, line, size=11)
            p.paragraph_format.space_after = Pt(10)
            index += 1
            continue
        if line.startswith("### Briefing"):
            briefing_count_in_account += 1
            if briefing_count_in_account > 1:
                doc.add_page_break()
            current_label = line.split()[-1].rstrip(".")
            doc.add_heading(f"Briefing {current_label}", level=2)
            index += 1
            continue
        if line == "#### Record scores before continuing":
            add_response_form(doc, current_label, include_pair=False)
            skip_generated_record_lines = True
            index += 1
            continue
        if line == "#### Record paired assessment":
            add_response_form(doc, current_label, include_pair=True)
            skip_generated_record_lines = True
            index += 1
            continue
        if line.startswith("#### "):
            heading_text = line[5:].strip()
            heading = doc.add_heading(heading_text, level=3)
            if heading_text == "Shared evidence register":
                heading.paragraph_format.page_break_before = True
            index += 1
            continue
        if line.startswith("|"):
            rows, index = parse_markdown_table(lines, index)
            add_table(doc, rows)
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            add_inline_markdown(p, line, size=9.5)
            style_paragraph(p, after=3)
            index += 1
            continue
        if line.startswith("-"):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, line[1:].strip(), size=9.5)
            style_paragraph(p, after=3)
            index += 1
            continue
        p = doc.add_paragraph()
        add_inline_markdown(p, line, size=9.5)
        style_paragraph(p, after=4)
        index += 1


def build_questionnaire(participant: str) -> Path:
    source = SOURCE_DIR / f"{participant}_participant_packet.md"
    doc = Document()
    configure_page(doc, participant)
    configure_styles(doc)
    add_cover(doc, participant)
    add_packet_content(doc, source)

    doc.core_properties.title = f"AccountLens Blinded Evaluation Questionnaire {participant}"
    doc.core_properties.subject = "Anonymous user evaluation questionnaire using fictional account data"
    doc.core_properties.author = "AccountLens course project"
    doc.core_properties.keywords = "blinded evaluation, synthetic data, account briefing"

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / f"AccountLens_Blinded_Questionnaire_{participant}.docx"
    doc.save(output)
    return output


def main() -> None:
    for participant in ("P01", "P02", "P03", "P04", "P05"):
        path = build_questionnaire(participant)
        print(f"Saved: {path}")


if __name__ == "__main__":
    main()
