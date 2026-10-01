from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
DOCX_OUTPUT = ROOT / "output" / "Wen_Hao_AccountLens_Final_Analysis.docx"
PDF_OUTPUT = ROOT / "output" / "pdf" / "Wen_Hao_AccountLens_Final_Analysis.pdf"

TITLE = "Enterprise Account Decision-Chain Mapping & Cross-Team Intelligence Aggregation Assistant"
SUBTITLE = "Final Project Analysis"
AUTHOR_LINE = "Wen Hao | PE6201 Emerging AI Technologies | 1 October 2026"

SECTIONS: list[tuple[str, list[str]]] = [
    (
        "1 Problem and significance",
        [
            "Enterprise account teams often hold customer intelligence in separate CRM notes, email threads, meeting records, calendars and support tickets. A Key Account Manager preparing for a client meeting may spend valuable time reconstructing who influences the decision, who can approve the purchase, what changed recently and which risk requires attention. Fragmentation can produce a briefing that is slow, incomplete or overconfident. AccountLens addresses this bounded preparation task: one account enters the workflow and one evidence-grounded Account Panorama Briefing comes out for human review.",
        ],
    ),
    (
        "2 Solution and scope",
        [
            "The prototype infers six decision-role classes, highlights a likely commercial signatory, maps evidence-linked workflow hypotheses, ranks activity from the previous 30 days, identifies unresolved tickets, unanswered email and competitor signals, and recommends no more than two next actions. It provides a local dashboard plus one-page PDF and Slack-ready Markdown exports. It does not connect to live Salesforce, Outlook, Jira or Slack, and it never sends a message or changes a source system. This keeps the project focused on the teacher's requested unit of value: one account and one meeting briefing.",
            "Under the Class 4 definition, AccountLens is a bounded AI-assisted workflow rather than a full agent. Application code fixes the sequence and performs filtering, deterministic risk checks, validation and export. The model performs one structured inference; it does not choose tools, vary the action sequence or receive observations in a thought-action loop. A full agent would add cost, compounded failure risk and governance exposure without improving the measured task.",
        ],
    ),
    (
        "3 Data and implementation",
        [
            "Because real enterprise systems were unavailable, the project uses reproducible fictional data. A fixed-seed generator creates 60 accounts for development, validation and frozen testing. A separate teacher-aligned set contains 10 hand-authored accounts. Each includes two email threads with explicit CC lists, one meeting attendee list, two support tickets, one CRM note, six contacts and a true_signatory_contact_id. The signatory is defined as the contact with final commercial approval authority. This label is removed before every model call and used only for offline scoring.",
            "The application rents a GPT-4o-compatible model and provider infrastructure while owning the data schema, prompts, orchestration, evidence validation, evaluation and governance rules. Exact account filtering is sufficient, so a vector database and RAG layer were not added. Structured output is checked against a Pydantic schema; invalid contact or evidence IDs trigger retry or visible failure. Twenty-four automated tests cover data generation, leakage prevention, role metrics, evidence validation, reporting and economics.",
        ],
    ),
    (
        "4 Evaluation method",
        [
            "The primary comparison is a deterministic rules baseline, B0, against the evidence-grounded model, M2. The 60-account dataset was split by account into 30 development, 10 validation and 20 test accounts. Model name, prompt and confidence threshold were frozen before one final test execution. Metrics include Macro-F1, accuracy, evidence precision, critical-event recall, abstention, p95 latency and estimated API cost. The supplementary 10-account evaluation measures signatory precision, recall and selection rate so that omissions cannot be hidden behind precision alone.",
            "Five anonymous participants also completed a counterbalanced blinded comparison of B0 and M2 briefings for four fictional validation accounts. They recorded preparation time, six five-point ratings and a preference. The study is reported descriptively because the sample is small, responses were self-reported and participant identity or live sessions were not independently verified.",
        ],
    ),
    (
        "5 Results",
        [
            "On the frozen 20-account test, M2 achieved 1.000 accuracy and Macro-F1, compared with B0 Macro-F1 of 0.7024. M2 evidence precision and critical-event recall were both 1.000, with p95 latency of 8.13 seconds. On the separate teacher-aligned set, B0 selected seven signatories and all seven were correct, giving 1.00 precision but only 0.70 recall and selection rate. M2 selected all 10 and matched all 10 labels, giving 1.00 precision, recall and selection rate. These are controlled fictional-data results, not proof of real-world accuracy.",
            "The frozen model run cost USD 0.321290 for 20 briefings, or USD 0.016065 per briefing. The teacher-aligned run cost USD 0.125933. In the blinded pilot, M2 was preferred in 20 of 20 paired comparisons, median preparation time fell from 185 to 145 seconds, and the median mean rating increased from 2.83 to 4.50. The uniform small sample makes the result directional rather than population-level evidence.",
        ],
    ),
    (
        "6 Business and technical trade-offs",
        [
            "The likely signatory is a Class 5 Archetype B output because the system manufactures a measurement that source systems do not directly record; briefing summarisation is an Archetype A subtask. The use case passes the prototype gates for worth, possibility, affordability, absorption and killability, but production affordability is unproven. The observed variable cost excludes integration, monitoring, maintenance, human review and governance. At the variable-only rate, a USD 10 balance supports about 622 briefings, but this is not a production budget.",
            "A future pilot should stop or be redesigned if signatory precision is below 0.85, selection rate below 0.70, evidence precision below 0.90, unsupported-claim rate above 0.05, p95 latency above 15 seconds, or full cost exceeds the measured value of preparation time saved.",
        ],
    ),
    (
        "7 Responsible use and limitations",
        [
            "The system is advisory. Confidence, abstention, evidence IDs, missing-role flags and human confirmation reduce the risk of redirecting account effort based on an unsupported role. Workflow edges are labelled as hypotheses rather than verified reporting lines. Source records are treated as untrusted data, raw content is hidden by default, logs contain metadata only, and no outbound write is available.",
            "The main limitation is synthetic data. The accounts are balanced and generated from repeated patterns, so the perfect M2 score does not establish performance across real organisations, languages, writing styles or data-quality failures. Production use would require lawful access, source provenance, independent signatory annotation, named data ownership, role-based access, retention controls, drift monitoring and a larger observed user study linked to downstream decision quality.",
        ],
    ),
    (
        "8 Conclusion",
        [
            "AccountLens demonstrates that a narrow, evidence-grounded workflow can turn fragmented account records into a faster and more actionable meeting briefing while preserving human authority. The contribution is not autonomy; it is a measurable combination of signatory selection, role coverage, evidence traceability, abstention, bounded recommendations and explicit failure handling. The prototype meets the course objective on fictional data and defines clear conditions that must be satisfied before any real-data pilot or production claim.",
        ],
    ),
]


def word_count() -> int:
    content = [TITLE, SUBTITLE, AUTHOR_LINE]
    for heading, paragraphs in SECTIONS:
        content.append(heading)
        content.extend(paragraphs)
    return len(re.findall(r"\b[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*\b", "\n".join(content)))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top: int = 85, start: int = 100, bottom: int = 85, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name: str, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def build_docx(count: int) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Heading 1", 12.5),):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run(TITLE), "Arial", 18, True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle.paragraph_format.keep_with_next = True
    set_run_font(subtitle.add_run(SUBTITLE), "Arial", 12.5, True)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(8)
    meta.paragraph_format.keep_with_next = True
    set_run_font(meta.add_run(f"{AUTHOR_LINE} | Word count: {count}"), "Arial", 9, False, "555555")

    for heading, paragraphs in SECTIONS:
        hp = doc.add_paragraph(style="Heading 1")
        hp.paragraph_format.space_before = Pt(7)
        hp.paragraph_format.space_after = Pt(3)
        hp.paragraph_format.keep_with_next = True
        set_run_font(hp.add_run(heading), "Arial", 12.5, True)
        for body in paragraphs:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.08
            set_run_font(p.add_run(body), "Arial", 10.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("Wen Hao | PE6201 Final Project Analysis"), "Arial", 8, False, "666666")

    core = doc.core_properties
    core.title = f"{TITLE} Final Project Analysis"
    core.subject = "PE6201 final project analysis"
    core.author = "Wen Hao"
    core.keywords = "AccountLens, PE6201, decision chain, signatory, evidence-grounded AI"
    core.comments = "Submission analysis capped below 1,200 words."

    DOCX_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUTPUT)


def build_pdf(count: int) -> None:
    PDF_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(PDF_OUTPUT),
        pagesize=LETTER,
        rightMargin=0.72 * inch,
        leftMargin=0.72 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.58 * inch,
        title=f"{TITLE} Final Project Analysis",
        author="Wen Hao",
        subject="PE6201 final project analysis",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AnalysisTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=21,
        textColor=colors.black,
        alignment=TA_LEFT,
        spaceAfter=5,
    )
    subtitle_style = ParagraphStyle(
        "AnalysisSubtitle",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12.5,
        leading=15,
        textColor=colors.black,
        alignment=TA_LEFT,
        spaceAfter=2,
    )
    meta_style = ParagraphStyle(
        "AnalysisMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#555555"),
        spaceAfter=8,
    )
    heading_style = ParagraphStyle(
        "AnalysisHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12.2,
        leading=14.5,
        textColor=colors.black,
        spaceBefore=6,
        spaceAfter=3,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "AnalysisBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.9,
        leading=12.2,
        textColor=colors.black,
        alignment=TA_LEFT,
        spaceAfter=5,
    )

    story = [
        Paragraph(TITLE.replace("&", "&amp;"), title_style),
        Paragraph(SUBTITLE, subtitle_style),
        Paragraph(f"{AUTHOR_LINE} | Word count: {count}", meta_style),
    ]
    for heading, paragraphs in SECTIONS:
        story.append(Paragraph(heading, heading_style))
        for body in paragraphs:
            story.append(Paragraph(body.replace("&", "&amp;"), body_style))

    def add_page_number(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawCentredString(LETTER[0] / 2, 0.32 * inch, f"{document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--count-only", action="store_true")
    args = parser.parse_args()
    count = word_count()
    print(f"word_count={count}")
    if count > 1200:
        raise SystemExit("Analysis exceeds 1,200 words")
    if args.count_only:
        return
    build_docx(count)
    build_pdf(count)
    print(DOCX_OUTPUT)
    print(PDF_OUTPUT)


if __name__ == "__main__":
    main()
