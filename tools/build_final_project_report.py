"""Build the submission-ready AccountLens final project report."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "AccountLens_Final_Project_Report.docx"

BLACK = "000000"
DARK_BLUE = "1F4E78"
LIGHT_BLUE = "EEF4F8"
LIGHT_GRAY = "F5F5F5"
BORDER = "D9D9D9"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), BORDER)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_repeat_table_header(table) -> None:
    if table.rows:
        repeat_header(table.rows[0])


def set_run_font(run, name="Aptos", size=11, bold=False, color=BLACK, italic=False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_run_font(run, size=9, color="666666")


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph_borders(paragraph_or_style) -> None:
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullet(doc: Document, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    run = p.add_run(text)
    set_run_font(run)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_numbered(doc: Document, number: int, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    number_run = p.add_run(f"{number}.  ")
    set_run_font(number_run)
    run = p.add_run(text)
    set_run_font(run)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    compact: bool = False,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell, top=60, bottom=60) if compact else set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, size=9.2 if compact else 9.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        body_row = table.add_row()
        prevent_row_split(body_row)
        cells = body_row.cells
        for i, value in enumerate(values):
            cell = cells[i]
            set_cell_margins(cell, top=55, bottom=55) if compact else set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i > 0 and len(value) < 24:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_run_font(run, size=9.0 if compact else 9.3)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    remove_paragraph_borders(title)

    for level, size in ((1, 17), (2, 13), (3, 11.5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def build() -> None:
    baseline = json.loads((ROOT / "results" / "baseline_test.json").read_text(encoding="utf-8"))
    m2 = json.loads((ROOT / "results" / "m2_test.json").read_text(encoding="utf-8"))
    user_study = json.loads((ROOT / "results" / "real_user_study.json").read_text(encoding="utf-8"))
    teacher_baseline = json.loads((ROOT / "results" / "teacher_aligned_baseline.json").read_text(encoding="utf-8"))
    teacher_m2 = json.loads((ROOT / "results" / "teacher_aligned_m2.json").read_text(encoding="utf-8"))
    economics = json.loads((ROOT / "results" / "class5_economics.json").read_text(encoding="utf-8"))

    b_metrics = baseline["metrics"]
    m_metrics = m2["metrics"]
    usage = m2["usage"]
    study = user_study["summary"]
    paired = study["paired_outcomes"]
    b0_signatory = teacher_baseline["signatory_metrics"]
    m2_signatory = teacher_m2["signatory_metrics"]
    teacher_usage = teacher_m2["usage"]

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Enterprise Account Decision-Chain Mapping & Cross-Team Intelligence Aggregation Assistant")
    remove_paragraph_borders(title)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("Final Project Report")
    set_run_font(run, name="Aptos Display", size=16, bold=True)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(24)
    set_run_font(meta.add_run("AccountLens  |  PE6201 End of Course Project  |  24 September 2026"), size=10.5, color="555555")

    add_heading(doc, "Project conclusion", 1)
    add_body(
        doc,
        "AccountLens demonstrates an end-to-end, evidence-grounded workflow for preparing an enterprise account meeting from fragmented fictional CRM, email, meeting, calendar, and support records. The selected M2 configuration achieved 1.000 accuracy and Macro-F1 on the frozen synthetic test split, with 1.000 evidence precision, 1.000 critical-event recall, and p95 latency of 8.1254 seconds. A separate instructor-aligned evaluation used 10 hand-authored accounts with explicit CC lists, attendee lists, two tickets and a true signatory per account; M2 selected all 10 signatories correctly. A five-person blinded pilot preferred M2 in all 20 paired comparisons and reduced median preparation time by 21.6%. These results support the prototype direction but do not establish production performance.",
    )

    add_heading(doc, "Report structure", 2)
    for item in [
        "Problem and significance",
        "Project scope and tradeoffs",
        "System design and workflow",
        "Data and evaluation method",
        "Results and user evidence",
        "Class 5 use case economics and decision gate",
        "Responsible use and safeguards",
        "Limitations and next development steps",
        "Demonstration and reproducibility",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()

    add_heading(doc, "1 Problem and significance", 1)
    add_body(
        doc,
        "Enterprise account teams often store relevant customer intelligence in separate systems and teams. A Key Account Manager preparing for a meeting may need to search CRM notes, email, meeting records, calendar items, and support tickets, then infer who influences the decision, what changed recently, which risks require attention, and what should happen next. This manual reconstruction is slow and can produce an incomplete or overly confident account view.",
    )
    add_body(
        doc,
        "The primary user is a Key Account Manager with less than 15 minutes to prepare. Sales Directors, Presales Engineers, and Customer Success Managers are secondary users of the same briefing. The intended change is practical: instead of inspecting every source record, the user reviews one Account Panorama Briefing, checks the evidence and uncertainty labels, confirms or corrects the decision-role hypotheses, and enters the meeting with a concise view of the buying workflow, recent activity, risks, and next actions.",
    )

    add_heading(doc, "Project objective", 2)
    add_body(
        doc,
        "The project objective is to test whether an evidence-grounded AI briefing can improve the usefulness and speed of enterprise meeting preparation relative to a deterministic rules baseline while preserving traceability, abstention, and human decision authority.",
    )
    add_table(
        doc,
        ["Success dimension", "Operational measure", "Project target"],
        [
            ["Role inference", "Macro-F1 and accuracy on a frozen account-level test split", "Improve on the non-AI baseline"],
            ["Commercial signatory", "Precision, recall, and selection rate on 10 instructor-aligned accounts", "Report omissions beside precision"],
            ["Traceability", "Evidence precision and critical-event recall", "Important claims remain auditable"],
            ["Speed", "p95 generation latency", "Under 15 seconds"],
            ["Meeting usefulness", "Blinded ratings, preference, and preparation time", "M2 improves ratings and reduces time"],
            ["Safety", "Abstention, uncertainty, human confirmation, and failure handling", "No unsupported automated action"],
        ],
        [1.35, 3.2, 2.3],
    )

    add_heading(doc, "2 Project scope and tradeoffs", 1)
    add_body(
        doc,
        "The minimum viable product uses one fictional enterprise account at a time and produces an Account Panorama Briefing. The briefing includes a likely commercial signatory, other decision roles, an evidence-linked workflow hypothesis, the most relevant activity from the previous 30 days, deterministic and model-supported risks, one or two next actions, evidence IDs, and limitations. The interface supports local one-page PDF and Slack-ready Markdown exports but never sends a message or changes a source system.",
    )
    add_table(
        doc,
        ["Decision", "Selected approach", "Reason"],
        [
            ["Model strategy", "Rent openai/gpt-4o through OpenRouter", "Avoid training cost while measuring a controlled structured-output task"],
            ["Application logic", "Build retrieval, validation, risks, mapping, evaluation, and exports", "Keep important controls inspectable and reproducible"],
            ["Retrieval", "Exact account filtering", "A vector database adds cost and failure modes without benefit for 60 synthetic accounts"],
            ["Control-flow design", "Fixed AI-assisted workflow, not a full agent", "Class 4 requires model-directed variable steps plus tool observations; this task needs neither"],
            ["Connectors", "Synthetic source records", "Avoid personal data, credentials, and production security exposure"],
            ["Slack", "Downloadable Markdown only", "Preserve human review before any external communication"],
        ],
        [1.25, 2.15, 3.45],
    )
    add_body(
        doc,
        "Live Salesforce, Outlook, Jira or Zendesk connections, production Slack posting, SSO, role-based access control, verified organisational reporting lines, and production deployment are outside the course-project scope. These are not hidden omissions: the interface and documentation state the boundaries directly.",
    )

    add_heading(doc, "3 System design and workflow", 1)
    add_body(
        doc,
        "AccountLens separates deterministic controls from model generation. This design keeps data selection, risk detection, evidence validation, failure handling, and export limits inspectable even when the model is unavailable or returns invalid output.",
    )
    for number, step in enumerate([
        "Select one fictional account and load its contacts and events.",
        "Filter events to the inclusive 30-day operating window and allow the user to exclude irrelevant items.",
        "Run the deterministic baseline, rank cross-team activity, and detect unresolved tickets, unanswered email, and competitor signals.",
        "Require the user to accept the responsible-use boundary before AI generation.",
        "Send structured, untrusted source records to the selected model with role definitions and evidence requirements.",
        "Validate the structured response, evidence IDs, confidence and abstention behaviour; retry or fail visibly if validation does not pass.",
        "Present the decision chain, missing roles, risks, actions, evidence register, limitations, and local exports for human confirmation.",
    ], start=1):
        add_numbered(doc, number, step)

    add_heading(doc, "Class 4 workflow classification", 2)
    add_body(
        doc,
        "Under the course definition, AccountLens is not a full agent. The model does not choose the next step, the step count is fixed, and no tool observation is fed back into a model-directed loop. The application therefore sits at the single-call rung inside deterministic orchestration. This deliberately avoids compounded step failures, unpredictable cost, and the governance cliff of an autonomous write tool.",
    )
    add_table(
        doc,
        ["Course test", "AccountLens", "Result"],
        [
            ["Who chooses the sequence", "Application code", "Workflow"],
            ["Variable runtime steps", "No", "Workflow"],
            ["Tool-observation loop", "None", "Not an agent"],
            ["External write authority", "None", "Advisory only"],
        ],
        [2.15, 2.95, 1.75],
        compact=True,
    )

    add_heading(doc, "Class 2 seven-layer product stack", 2)
    add_table(
        doc,
        ["Layer", "Implementation", "Decision"],
        [
            ["Compute", "API-provider infrastructure", "Rent"],
            ["Data and embeddings", "Owned fictional records and labels; embeddings unnecessary", "Build and own"],
            ["Vector store", "None; exact account filtering", "Do not add"],
            ["Model", "Pinned GPT-4o-compatible endpoint", "Rent"],
            ["Orchestration", "Filtering, validation, risk rules and export", "Build and own"],
            ["Serving", "Local Streamlit demonstration", "Build"],
            ["Observability and evals", "Logs, frozen splits, 24 tests and human pilot", "Build and own"],
        ],
        [1.6, 3.8, 1.45],
        compact=True,
    )
    add_body(
        doc,
        "This follows the course build-versus-buy rule: own the data, orchestration, evaluation and governance while renting commodity compute and the model. Retrieval-augmented generation is not used because the records are already structured and selected by exact account ID; adding embeddings or a vector database would add cost and failure modes without improving this measured task.",
    )

    add_heading(doc, "Core decision roles", 2)
    add_table(
        doc,
        ["Role class", "Operational meaning"],
        [
            ["Champion", "Coordinates internal support, advocates for the solution, and drives adoption"],
            ["Economic buyer", "Controls final commercial approval and is evaluated as the likely signatory"],
            ["Technical evaluator", "Assesses architecture, security, integration, or standards"],
            ["Procurement legal", "Manages commercial terms, supplier onboarding, or legal review"],
            ["End user", "Represents operational use, workflow, or adoption needs"],
            ["Unknown", "Evidence is insufficient for a supported role assignment"],
        ],
        [1.65, 5.2],
        compact=True,
    )

    add_heading(doc, "4 Data and evaluation method", 1)
    add_body(
        doc,
        "A fixed-seed generator created 60 fictional enterprise accounts with contacts, role ground truth, varied CRM, email, meeting, calendar, and support events, and one true_signatory_contact_id per account. The true signatory is the contact with final commercial approval authority. The field is excluded from model input and is used only for offline evaluation. Account-level splits were frozen before experiments: 30 development accounts, 10 validation accounts, and 20 test accounts.",
    )
    add_body(
        doc,
        "A separate teacher-aligned set contains 10 hand-authored fictional accounts. Each account has two email threads with explicit CC lists, one meeting attendee list, two support tickets, one CRM note, six contacts and one true commercial signatory. Its manifest records counts and file hashes. This supplementary set is isolated from the frozen 60-account experiment, and the signatory label is removed before every model call.",
    )
    add_table(
        doc,
        ["Split", "Accounts", "Use"],
        [
            ["Development", "30", "Implementation, prompt iteration, and smoke testing"],
            ["Validation", "10", "Experiment comparison and threshold selection"],
            ["Test", "20", "One-time final evaluation after freezing configuration"],
        ],
        [1.4, 1.2, 4.25],
    )
    add_body(
        doc,
        "Four configurations were compared: B0 deterministic rules, M1 direct-model baseline, M2 evidence-grounded model, and H1 hybrid rules plus model. M2 was selected because it matched H1's validation quality with lower estimated cost and lower p95 latency while retaining evidence requirements and abstention behaviour. The final threshold of 0.65 was chosen on validation data; thresholds from 0.50 to 0.85 tied, so the preset value was retained to avoid unnecessary overfitting.",
    )
    add_body(
        doc,
        "Before the final run, the model name, prompt version, threshold, and pipeline were frozen. The M2 test split was then run once. The saved model results were not regenerated after review. Post-evaluation interface features use the frozen results and do not alter the reported classification metrics.",
    )

    add_heading(doc, "Blinded user study method", 2)
    add_body(
        doc,
        "Five anonymous participants reviewed B0 and M2 briefings for four validation accounts in counterbalanced A/B order. They recorded preparation time, meeting readiness, six ratings on a five-point scale, a paired preference, and qualitative comments. The A/B key was decoded only after completed questionnaires were retrieved. The resulting 20 paired comparisons are reported descriptively because the sample is small.",
    )

    add_heading(doc, "5 Results and user evidence", 1)
    add_heading(doc, "Frozen test results", 2)
    add_table(
        doc,
        ["Metric", "B0 rules", "Final M2"],
        [
            ["Accuracy", f"{b_metrics['accuracy']:.4f}", f"{m_metrics['accuracy']:.4f}"],
            ["Macro-F1", f"{b_metrics['macro_f1']:.4f}", f"{m_metrics['macro_f1']:.4f}"],
            ["Coverage", f"{b_metrics['coverage']:.4f}", f"{m_metrics['coverage']:.4f}"],
            ["Evidence precision", "N/A", f"{m2['evidence_precision']:.4f}"],
            ["Critical-event recall", "N/A", f"{m2['critical_event_recall']:.4f}"],
            ["Estimated API cost", "$0.000000", f"${usage['estimated_cost_usd']:.6f}"],
            ["p95 latency", "N/A", f"{usage['p95_latency_seconds']:.4f} s"],
        ],
        [3.2, 1.8, 1.8],
    )
    add_body(
        doc,
        f"The test split contains {m_metrics['samples']} role labels across 20 accounts. All 20 M2 calls succeeded on the first attempt. M2 used {usage['input_tokens']:,} input tokens and {usage['output_tokens']:,} output tokens. Its 0.8333 coverage is intentional: each account includes one correctly identified unknown contact that is counted as an abstention. The rules baseline made 40 errors by abstaining on supported non-unknown roles.",
    )

    add_heading(doc, "Commercial signatory evaluation", 2)
    add_table(
        doc,
        ["System", "Selected", "Correct", "Precision", "Recall", "Selection rate"],
        [
            ["B0 rules", str(b0_signatory["selections"]), str(b0_signatory["correct_selections"]), f"{b0_signatory['precision']:.4f}", f"{b0_signatory['recall']:.4f}", f"{b0_signatory['selection_rate']:.4f}"],
            ["M2", str(m2_signatory["selections"]), str(m2_signatory["correct_selections"]), f"{m2_signatory['precision']:.4f}", f"{m2_signatory['recall']:.4f}", f"{m2_signatory['selection_rate']:.4f}"],
        ],
        [1.3, 0.9, 0.9, 1.15, 1.15, 1.35],
        compact=True,
    )
    add_body(
        doc,
        f"This supplementary evaluation used the separate 10-account instructor-aligned set. Precision alone would make B0 and M2 appear equal, but B0 omitted three true signatories. M2 selected one signatory for every account and matched all 10 fictional ground-truth labels. The ten M2 calls used {teacher_usage['input_tokens']:,} input tokens and {teacher_usage['output_tokens']:,} output tokens, cost an estimated ${teacher_usage['estimated_cost_usd']:.6f}, and had p95 latency of {teacher_usage['p95_latency_seconds']:.4f} seconds. All evidence validations passed. The result does not establish real-world signatory accuracy.",
    )

    add_heading(doc, "Blinded pilot results", 2)
    b0_study = study["systems"]["B0"]
    m2_study = study["systems"]["M2"]
    add_table(
        doc,
        ["Measure", "B0 median", "M2 median"],
        [
            ["Role usefulness", "2.0", "5.0"],
            ["Evidence trust", "4.0", "5.0"],
            ["Risk relevance", "2.0", "4.0"],
            ["Action usefulness", "2.0", "5.0"],
            ["Clarity", "3.0", "4.0"],
            ["Uncertainty", "4.0", "4.0"],
            ["Mean of six ratings", f"{b0_study['median_mean_rating']:.2f}", f"{m2_study['median_mean_rating']:.2f}"],
            ["Preparation time", f"{b0_study['median_preparation_time_seconds']:.0f} s", f"{m2_study['median_preparation_time_seconds']:.0f} s"],
        ],
        [3.2, 1.8, 1.8],
    )
    add_body(
        doc,
        f"M2's median mean rating was {paired['median_mean_rating_gain_m2_minus_b0']:.2f} points higher. Median preparation time was {paired['median_preparation_time_reduction_percent']:.1f}% lower, and M2 was preferred in {paired['m2_preferences']}/{paired['total_pairs']} comparisons ({paired['m2_preference_percent']:.0f}%). All three pre-registered pilot targets were met.",
    )
    add_body(
        doc,
        "Participants repeatedly described M2 as more meeting-ready because it names likely decision roles, links claims to evidence, and makes the next action easier to identify. Their recurring caution was that role authority remains a hypothesis and should be verified before a customer conversation. B0 was viewed as a summary of activity rather than an explicit map of the buying group.",
    )

    add_heading(doc, "Interpretation", 2)
    add_body(
        doc,
        "Together, the two evaluations support different claims. The frozen synthetic test shows that M2 can recover the generator's role patterns with complete evidence linkage under controlled conditions. The blinded pilot suggests that the same evidence-grounded structure can reduce meeting-preparation effort and improve perceived actionability. Neither result proves performance on real enterprise data, real organisational power structures, or downstream revenue outcomes.",
    )

    add_heading(doc, "6 Class 5 use case economics and decision gate", 1)
    add_heading(doc, "AI opportunity archetype", 2)
    add_body(
        doc,
        "The core signatory and decision-chain task is Archetype B because it manufactures a measurement that is not recorded directly in source systems. Observable proxies include budget discussions, meeting attendance, email participation, procurement activity, tickets, and CRM notes. The synthetic ground truth acts as a calibration set. Briefing summarisation is an Archetype A subtask because it replaces manual account reconstruction. The main Archetype B risk is that an estimate is read as a fact, so the interface labels the signatory as likely and requires evidence review and human confirmation.",
    )

    add_heading(doc, "Use case intake screen", 2)
    add_table(
        doc,
        ["Gate", "Prototype evidence", "Decision"],
        [
            ["Worth doing", "Repeated meeting-preparation task; 40-second directional median time difference; scale value, limited scope value, no free learning label", "Pilot pass"],
            ["Possible", "Synthetic data pass the five readiness tests; real source access and consistency are not proven", "Prototype pass only"],
            ["Affordable", "Observed variable cost is USD 0.016065; human fallback and fixed costs remain assumptions", "Production pending"],
            ["Absorbable", "KAM reviews evidence and verifies the signatory before a meeting; no outbound or write action", "Advisory pass"],
            ["Killable", "Original thresholds plus future signatory, full-cost, and data-ownership gates", "Pilot pass"],
        ],
        [1.15, 4.35, 1.35],
        compact=True,
    )

    add_heading(doc, "Data readiness and ownership", 2)
    add_table(
        doc,
        ["Test", "Synthetic prototype", "Production requirement"],
        [
            ["Long", "Designed variation across 60 accounts", "History across organisations, stages, writing styles, and outcomes"],
            ["Standardised", "One schema and role definition", "Reconciled source fields and signatory definition"],
            ["Quality controlled", "Fixed truth and deterministic checks", "Named data steward and correction workflow"],
            ["Consistent", "Stable generated definitions", "Data contracts and drift monitoring"],
            ["Accessible", "Local fictional records", "Lawful permission, provenance, retention, and source ownership"],
        ],
        [1.15, 2.35, 3.35],
        compact=True,
    )
    add_body(
        doc,
        "The weakest production links are lawful access, cross-team consistency, daily data ownership, and independently verified signatory labels. The KAM owns the briefing decision. Sales Operations or a named data steward owns definitions and corrections. Source-system owners, Security, and Legal approve access and governance.",
    )

    add_heading(doc, "Cost per successful briefing", 2)
    add_body(
        doc,
        f"The frozen 20-account test cost USD {economics['observed_test_cost_usd']:.6f}, or USD {economics['observed_variable_cost_per_briefing_usd']:.6f} per briefing. At that variable-only rate, USD 10 supports approximately {economics['ten_dollar_capacity_at_observed_variable_cost']} briefings. Full cost to serve equals variable cost plus expected human fallback cost plus fixed monthly cost divided by monthly volume.",
    )
    scenario_rows = []
    for scenario in economics["scenarios"]:
        scenario_rows.append(
            [
                scenario["name"],
                f"{scenario['success_rate']:.0%}",
                f"${scenario['expected_fallback_cost_usd']:.3f}",
                f"${scenario['fixed_cost_per_task_usd']:.3f}",
                f"${scenario['cost_per_successful_briefing_usd']:.3f}",
                f"${scenario['break_even_hourly_rate_usd']:.2f}",
            ]
        )
    add_table(
        doc,
        ["Scenario", "Success", "Fallback", "Fixed per task", "Total", "Break-even wage"],
        scenario_rows,
        [2.0, 0.75, 0.95, 1.1, 0.85, 1.2],
        compact=True,
    )
    add_body(
        doc,
        "Only the first row is an observed project estimate. The other rows are sensitivity assumptions. Production affordability cannot be claimed until real success rate, review time, loaded labour cost, monthly volume, integration, monitoring, maintenance, and governance costs are measured.",
    )

    add_heading(doc, "Stop conditions", 2)
    for condition in [
        "Stop or re-scope a real-data pilot if signatory precision is below 0.85 or selection rate is below 0.70 on an independently labelled calibration set.",
        "Stop if evidence precision is below 0.90, unsupported-claim rate exceeds 0.05, or p95 generation latency exceeds 15 seconds.",
        "Stop if cost per successful briefing exceeds the measured value of preparation time saved.",
        "Do not start a production pilot without lawful access, provenance, daily data ownership, and a correction workflow.",
    ]:
        add_bullet(doc, condition)

    add_heading(doc, "7 Responsible use and safeguards", 1)
    add_table(
        doc,
        ["Risk", "Implemented mitigation"],
        [
            ["Incorrect role inference redirects account effort", "Confidence scores, abstention, evidence IDs, missing-role flags, and mandatory human confirmation"],
            ["Workflow hypothesis is mistaken for a reporting line", "Edges are labelled as workflow hypotheses and repeated limitations require verification"],
            ["Prompt injection in source records", "All account records are treated as untrusted data and adversarial examples are included in testing"],
            ["Unsupported model output", "Structured schema, claim-level evidence validation, retry handling, and visible failure states"],
            ["Sensitive data exposure", "Fictional data, hidden raw content by default, metadata-only logs, and no autonomous outbound action"],
            ["Over-reliance or action without review", "Advisory-only interface; export and generation require a human confirmation boundary"],
            ["Information overload", "Ranked activity, a 30-day window, concise risks, and no more than two recommended actions"],
        ],
        [2.35, 4.5],
    )
    add_body(
        doc,
        "The application is designed as decision support, not a decision maker. It does not verify identity, authority, employment status, reporting relationships, or customer consent. Users remain responsible for checking source evidence, correcting role hypotheses, following organisational data-governance rules, and deciding whether any recommendation is appropriate.",
    )

    add_heading(doc, "8 Limitations and next development steps", 1)
    add_heading(doc, "Current limitations", 2)
    for limitation in [
        "The dataset is synthetic, balanced, and generated from repeated patterns, making the task easier than noisy enterprise data.",
        "The perfect M2 test result measures generalisation across fictional accounts created by the same generator, not across organisations, languages, CRM systems, or real data-quality failures.",
        "The blinded pilot contains five self-reported participants and highly uniform answers. Questionnaire completeness and arithmetic were verified, but participant identity and live sessions were not independently observed.",
        "Preparation time is self-recorded and does not measure meeting quality, sales-cycle progression, customer satisfaction, or revenue impact.",
        "Latency and estimated cost depend on the external model provider and are not guaranteed production service levels.",
        "The prototype lacks production authentication, connector permissions, retention controls, monitoring, and formal source-correction workflows.",
    ]:
        add_bullet(doc, limitation)

    add_heading(doc, "Recommended next development steps", 2)
    for number, step in enumerate([
        "Repeat evaluation with privacy-approved, de-identified enterprise examples and independent role annotation.",
        "Run a larger, externally observed user study with more varied account-management and technical roles.",
        "Measure decision quality and follow-up outcomes in addition to short-term preparation time.",
        "Add role-based access control, audit logging, retention rules, security review, and read-only connector permissions.",
        "Monitor confidence calibration, distribution shift, abstention quality, connector failures, and evidence-validation failures.",
    ], start=1):
        add_numbered(doc, number, step)

    add_heading(doc, "9 Demonstration and reproducibility", 1)
    add_body(
        doc,
        "The repository contains the fixed-seed generator, instructor-aligned ten-account dataset, account-level data-split logic, deterministic baseline, model pipeline, signatory evaluator, cost-to-serve calculator, evidence validator, test runner, Streamlit interface, export layer, study materials, normalized study results, and evaluation documentation. Twenty-four automated tests currently pass. A clean-environment verification confirms that the documented setup and application startup work from a fresh environment.",
    )
    add_heading(doc, "Suggested demonstration sequence", 2)
    for number, step in enumerate([
        "Show the original project title, fictional-data notice, one selected account, and the fixed 30-day window.",
        "Inspect the decision chain, missing roles, cross-team activity, deterministic risks, and evidence IDs.",
        "Confirm the responsible-use boundary, generate a briefing, and show one or two next actions and limitations.",
        "Demonstrate calibrated abstention for an ambiguous contact rather than forcing a role.",
        "Demonstrate a handled API failure in a separate temporary session without exposing the saved key.",
        "Present frozen test metrics and the five-person pilot result, then close with limitations and next steps.",
    ], start=1):
        add_numbered(doc, number, step)

    add_heading(doc, "Evidence and artifact map", 2)
    add_table(
        doc,
        ["Evidence", "Repository location"],
        [
            ["Project setup and use", "README.md"],
            ["Problem statement alignment", "docs/statement_alignment.md"],
            ["Classes 1 to 5 alignment", "docs/course_alignment_classes_1_to_5.md"],
            ["Class 5 business case", "docs/class5_business_case.md"],
            ["Validation and threshold selection", "docs/phase2_results.md"],
            ["Frozen evaluation and integrity", "docs/final_evaluation.md"],
            ["Blinded pilot and limitations", "docs/real_user_evaluation.md"],
            ["Machine-readable pilot results", "results/real_user_study.json and real_user_study CSV files"],
            ["Test artifacts", "results/baseline_test.json and results/m2_test.json"],
            ["Signatory and economics artifacts", "results/teacher_aligned_m2.json and results/class5_economics.json"],
            ["Instructor-aligned dataset", "data/teacher_aligned_10 and docs/teacher_aligned_dataset.md"],
            ["Instructor-aligned results", "results/teacher_aligned_baseline.json and results/teacher_aligned_m2.json"],
            ["Demonstration procedure", "docs/demo_script.md"],
            ["Submission status", "docs/submission_checklist.md"],
        ],
        [2.55, 4.3],
    )

    add_heading(doc, "10 Conclusion", 1)
    add_body(
        doc,
        "AccountLens fulfils the proposed project direction at prototype level: it selects a likely commercial signatory, maps other decision roles, aggregates cross-team intelligence from the previous 30 days, identifies risks, recommends limited next actions, and exports a concise briefing while preserving traceable evidence and human control. It is deliberately a bounded AI-assisted workflow rather than a full agent. The non-AI baseline, frozen experiment process, separate signatory evaluation, cost-to-serve analysis, small blinded pilot, and explicit limitations make the contribution measurable without presenting synthetic results as production proof.",
    )
    add_body(
        doc,
        "The most important result is not the perfect synthetic classification score alone. It is the combination of improved role coverage, evidence validation, calibrated abstention, faster meeting preparation in the pilot, visible failure handling, and a clear boundary between AI advice and accountable human action.",
    )

    add_heading(doc, "Appendix Evaluation integrity record", 1)
    add_table(
        doc,
        ["Item", "Recorded value"],
        [
            ["Selected experiment", "M2 evidence-grounded briefing"],
            ["Model and prompt", "openai/gpt-4o with account_briefing_v2.txt"],
            ["Frozen threshold", "0.65"],
            ["Frozen test execution", "One run on 20 accounts; no post-result rerun"],
            ["Prompt SHA-256", "A0BE5A8F89381857F2F12671C447812316D227FCD79EFEEE906C5921DBAAD541"],
            ["Pipeline SHA-256", "4733D34B32FCE317E9EAF860C40C75969A606E0C53B3498927073BCA90ED2A04"],
            ["Blinded pilot integrity", "5 complete questionnaires, 40 observations, 20 paired comparisons"],
            ["Signatory evaluation", "10 instructor-aligned accounts; one frozen M2 run; 10 of 10 correct"],
            ["Ground-truth leakage control", "true_signatory_contact_id is removed from every model prompt"],
            ["Automated verification", "24 project tests passed"],
        ],
        [2.15, 4.7],
    )
    add_body(
        doc,
        "The five questionnaire files were archived with matching source and archive SHA-256 checksums. The import did not rerun the frozen test split. Identity and live-session verification were outside the import process and are therefore not claimed.",
    )

    add_heading(doc, "References", 1)
    for reference in [
        "PE6201 Project Problem Statement Enterprise Account Decision-Chain Mapping & Cross-Team Intelligence Aggregation Assistant.",
        "PE6201 Project Proposal Watchouts.",
        "PE6201 Classes 1 to 4 course materials covering AI task selection, the modern AI stack, structured prompting and evaluation, workflow-versus-agent classification, hardening, and token economics.",
        "PE6201 Class 5 Where We Are, Where Is AI Worth Building, What Does It Actually Cost, and Two Archetypes and the Ant Group Case.",
        "AccountLens project documentation dated 24 September 2026.",
        "Frozen test artifacts and the real blinded user-evaluation evidence.",
    ]:
        add_bullet(doc, reference)

    core = doc.core_properties
    core.title = "Enterprise Account Decision-Chain Mapping & Cross-Team Intelligence Aggregation Assistant Final Project Report"
    core.subject = "AccountLens PE6201 final project report"
    core.author = "AccountLens project"
    core.keywords = "AccountLens, enterprise account, decision chain, evidence grounded AI, PE6201"
    core.comments = "Generated from the verified project artifacts on 24 September 2026."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
